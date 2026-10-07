"""Convert each supported input into source-located text and optional image bytes."""
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
import json
import statistics
import subprocess
import tempfile

from app.core.config import get_settings


@dataclass
class Segment:
    text: str
    location: str
    modality: str = "text"
    image: bytes | None = None
    preserve: bool = False
    locator: dict = field(default_factory=dict)
    quality: dict = field(default_factory=dict)


def ocr(image, lang: str | None = None) -> str:
    from PIL import Image
    if isinstance(image, bytes):
        image = Image.open(BytesIO(image))
    cfg = get_settings()
    if cfg.ocr_engine == "paddle":
        if any(language != "eng" for language in (lang or cfg.ocr_lang).split("+")):
            raise ValueError("This PaddleOCR configuration supports English only; use Tesseract with Hindi/Telugu packs")
        import numpy as np
        from paddleocr import PaddleOCR
        if not hasattr(ocr, "_paddle"):
            ocr._paddle = PaddleOCR(lang="en", use_angle_cls=True)
        result = ocr._paddle.ocr(np.asarray(image.convert("RGB")), cls=True)
        return "\n".join(line[1][0] for page in result or [] for line in page or [])
    if cfg.ocr_engine != "tesseract": raise ValueError("Unsupported OCR_ENGINE")
    import pytesseract
    return pytesseract.image_to_string(image, lang=lang or cfg.ocr_lang).strip()


def ocr_details(image: bytes, lang: str | None = None) -> tuple[str, float | None, list[int] | None]:
    """Return OCR text plus an uncalibrated engine signal and its pixel bounds."""
    cfg = get_settings()
    if cfg.ocr_engine != "tesseract":
        return ocr(image, lang), None, None
    from PIL import Image
    import pytesseract
    picture = Image.open(BytesIO(image)).convert("RGB")
    data = pytesseract.image_to_data(picture, lang=lang or cfg.ocr_lang, output_type=pytesseract.Output.DICT)
    words, scores, bounds = [], [], []
    for i, raw in enumerate(data["text"]):
        word = raw.strip()
        if not word: continue
        words.append(word)
        try:
            score = float(data["conf"][i])
            if score >= 0: scores.append(score)
        except (TypeError, ValueError):
            pass
        left, top, width, height = (int(data[key][i]) for key in ("left", "top", "width", "height"))
        bounds.append([left, top, left + width, top + height])
    box = [min(b[0] for b in bounds), min(b[1] for b in bounds),
           max(b[2] for b in bounds), max(b[3] for b in bounds)] if bounds else None
    return " ".join(words), round(sum(scores) / (100 * len(scores)), 3) if scores else None, box


def caption(image: bytes) -> str:
    if not get_settings().enable_image_caption:
        return ""
    from PIL import Image
    from transformers import pipeline
    # Opt-in downloads a caption model on first use.
    if not hasattr(caption, "_model"):
        caption._model = pipeline("image-to-text", model="Salesforce/blip-image-captioning-base")
    return caption._model(Image.open(BytesIO(image)))[0]["generated_text"]


def visual(image: bytes, location: str, locator: dict | None = None, lang: str | None = None) -> Segment:
    labels, signal, box = ocr_details(image, lang)
    description = caption(image)
    text = f"Image at {location}. " + (f"Visible text and chart labels: {labels}. " if labels else "") + (f"Visual description: {description}." if description else "")
    source_locator = dict(locator or {"kind": "image"})
    if box: source_locator["ocr_bbox_pixels"] = box
    return Segment(text.strip(), location, "image", image, locator=source_locator,
                   quality={"method": "ocr+caption" if description else "ocr", "ocr_signal": signal,
                            "flags": ["ocr"] + (["generated_caption"] if description else []),
                            "note": "OCR signal is not a calibrated probability; chart values require source review."})


def _log_insights(rows: list[dict]) -> str:
    if not rows:
        return ""
    errors = sum(any(str(v).lower() in {"error", "critical", "failed"} for v in row.values()) for row in rows)
    numeric = {}
    for row in rows:
        for k, v in row.items():
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                numeric.setdefault(str(k), []).append(float(v))
    findings = [f"Records: {len(rows)}; rows with error/critical/failed values: {errors}."]
    for key, values in list(numeric.items())[:10]:
        if len(values) < 3:
            continue
        median = statistics.median(values)
        mad = statistics.median(abs(v - median) for v in values)
        outliers = sum(abs(v - median) > 3 * 1.4826 * mad for v in values) if mad else 0
        findings.append(f"{key}: median {median:.3g}, min {min(values):.3g}, max {max(values):.3g}, robust outliers {outliers}.")
    return "Log pattern summary: " + " ".join(findings)


def extract(path: Path, extension: str, ocr_languages: str | None = None) -> list[Segment]:
    if extension in {".txt", ".log"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        return [Segment(text, "text", locator={"kind": "text"}, quality={"method": "native_text", "flags": []})]
    if extension in {".json", ".jsonl"}:
        raw = path.read_text(encoding="utf-8", errors="replace")
        if extension == ".jsonl":
            data = [json.loads(line) for line in raw.splitlines() if line.strip()]
        else:
            data = json.loads(raw)
        rows = data if isinstance(data, list) else data.get("records", []) if isinstance(data, dict) else []
        rows = [r for r in rows if isinstance(r, dict)]
        return [Segment(json.dumps(data, ensure_ascii=False, indent=2), "JSON", "logs", preserve=True,
                        locator={"kind": "json", "pointer": "/"}, quality={"method": "parsed_json", "flags": []}),
                Segment(_log_insights(rows), "log analysis", "logs", locator={"kind": "json", "pointer": "/records"},
                        quality={"method": "derived_summary", "flags": ["derived_analysis"]})]
    if extension == ".pdf":
        import fitz
        import pdfplumber
        result = []
        with fitz.open(path) as document, pdfplumber.open(path) as pdf:
            for i, page in enumerate(document):
                loc = f"page {i+1}"
                text = page.get_text(sort=True).strip()
                if text:
                    result.append(Segment(text, loc, locator={"kind": "pdf", "page": i+1, "anchor": text[:100]},
                                          quality={"method": "native_text", "flags": []}))
                tables = pdf.pages[i].extract_tables()
                if get_settings().pdf_table_engine == "camelot":
                    import camelot
                    tables = [t.df.values.tolist() for t in camelot.read_pdf(str(path), pages=str(i+1), flavor="stream")]
                for table_index, t in enumerate(tables):
                    lines = [" | ".join(str(c or "").replace("\n", " ") for c in row) for row in t if row]
                    if lines:
                        result.append(Segment("\n".join(lines), f"{loc} table {table_index+1}", "table", preserve=True,
                                              locator={"kind": "pdf", "page": i+1, "table": table_index+1},
                                              quality={"method": "parsed_table", "flags": ["table_extraction"]}))
                if len(text) < 40:
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    image = pix.tobytes("png")
                    result.append(visual(image, loc + " scanned page", {"kind": "pdf", "page": i+1, "render_scale": 2}, ocr_languages))
                else:
                    for j, info in enumerate(page.get_images(full=True)[:10]):
                        image = document.extract_image(info[0])["image"]
                        try:
                            result.append(visual(image, f"{loc} image {j+1}", {"kind": "pdf", "page": i+1, "image": j+1}, ocr_languages))
                        except Exception:
                            continue
        if not result and get_settings().unstructured_fallback:
            from unstructured.partition.auto import partition
            result = [Segment(str(element), "unstructured") for element in partition(filename=str(path))]
        return result
    if extension == ".docx":
        from docx import Document
        doc = Document(path)
        result = [Segment(p.text, f"paragraph {i+1}", locator={"kind": "docx", "paragraph": i+1},
                          quality={"method": "native_text", "flags": []}) for i, p in enumerate(doc.paragraphs) if p.text.strip()]
        for i, table in enumerate(doc.tables):
            rows = [" | ".join(cell.text.replace("\n", " ") for cell in row.cells) for row in table.rows]
            result.append(Segment("\n".join(rows), f"table {i+1}", "table", preserve=True,
                                  locator={"kind": "docx", "table": i+1}, quality={"method": "parsed_table", "flags": []}))
        if not result and get_settings().unstructured_fallback:
            from unstructured.partition.auto import partition
            result = [Segment(str(element), "unstructured") for element in partition(filename=str(path))]
        return result
    if extension == ".pptx":
        from pptx import Presentation
        result = []
        for i, slide in enumerate(Presentation(path).slides):
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    result.append(Segment(shape.text, f"slide {i+1}", locator={"kind": "pptx", "slide": i+1},
                                          quality={"method": "native_text", "flags": []}))
                if shape.has_table:
                    rows = [" | ".join(cell.text for cell in row.cells) for row in shape.table.rows]
                    result.append(Segment("\n".join(rows), f"slide {i+1} table", "table", preserve=True,
                                          locator={"kind": "pptx", "slide": i+1, "table": True},
                                          quality={"method": "parsed_table", "flags": []}))
                if shape.shape_type == 13:
                    try:
                        result.append(visual(shape.image.blob, f"slide {i+1} image", {"kind": "pptx", "slide": i+1, "image": True}, ocr_languages))
                    except Exception:
                        continue
        return result
    if extension in {".csv", ".xlsx"}:
        import pandas as pd
        sheets = pd.read_excel(path, sheet_name=None, nrows=100000) if extension == ".xlsx" else {"CSV": pd.read_csv(path, nrows=100000)}
        result = []
        for sheet, frame in sheets.items():
            for start in range(0, len(frame), 30):
                result.append(Segment(frame.iloc[start:start+30].fillna("").to_csv(index=False), f"sheet {sheet} rows {start+1}-{min(len(frame), start+30)}", "table", preserve=True,
                                      locator={"kind": "spreadsheet", "sheet": sheet, "row_start": start+1,
                                               "row_end": min(len(frame), start+30), "columns": [str(c) for c in frame.columns]},
                                      quality={"method": "parsed_table", "flags": []}))
            numeric = frame.select_dtypes(include="number")
            if not numeric.empty:
                result.append(Segment("Numeric summary:\n" + numeric.describe().to_string(), f"sheet {sheet} summary", "table", preserve=True,
                                      locator={"kind": "spreadsheet", "sheet": sheet, "summary": True},
                                      quality={"method": "derived_summary", "flags": ["derived_analysis"]}))
        return result
    if extension in {".png", ".jpg", ".jpeg", ".webp"}:
        return [visual(path.read_bytes(), "image", {"kind": "image"}, ocr_languages)]
    if extension in {".mp3", ".wav", ".m4a"}:
        from faster_whisper import WhisperModel
        if not hasattr(extract, "_whisper"):
            extract._whisper = WhisperModel(get_settings().whisper_model, compute_type="int8")
        segments, _ = extract._whisper.transcribe(str(path), vad_filter=True)
        return [Segment(s.text, f"{s.start:.1f}-{s.end:.1f}s", "audio",
                        locator={"kind": "audio", "start_seconds": round(s.start, 2), "end_seconds": round(s.end, 2)},
                        quality={"method": "asr", "asr_logprob": round(s.avg_logprob, 3), "flags": ["asr"],
                                 "note": "ASR log probability is not calibrated confidence."}) for s in segments]
    if extension in {".mp4", ".mov"}:
        import cv2
        cap = cv2.VideoCapture(str(path))
        if not cap.isOpened():
            raise ValueError("Could not decode video")
        fps = cap.get(cv2.CAP_PROP_FPS) or 25
        frame_count = cap.get(cv2.CAP_PROP_FRAME_COUNT)
        if frame_count / fps > 3600:
            cap.release()
            raise ValueError("Video limit is one hour")
        result = []
        step = max(1, int(fps * 15))
        index = 0
        previous = None
        try:
            while True:
                cap.set(cv2.CAP_PROP_POS_FRAMES, index)
                ok, frame = cap.read()
                if not ok:
                    break
                small = cv2.resize(frame, (64, 64))
                diff = cv2.norm(small, previous, cv2.NORM_L2) / small.size if previous is not None else 999
                previous = small
                if diff > 15 or not result:
                    _, encoded = cv2.imencode(".png", frame)
                    result.append(visual(encoded.tobytes(), f"video {index/fps:.1f}s",
                                         {"kind": "video", "second": round(index/fps, 2)}, ocr_languages))
                index += step
                if len(result) >= 240:
                    break
        finally:
            cap.release()
        with tempfile.TemporaryDirectory() as temp:
            audio = Path(temp) / "audio.wav"
            run = subprocess.run(["ffmpeg", "-nostdin", "-loglevel", "error", "-i", str(path), "-vn", "-ac", "1", "-ar", "16000", str(audio)], capture_output=True, timeout=120, check=False)
            if run.returncode == 0 and audio.exists():
                for segment in extract(audio, ".wav", ocr_languages):
                    segment.location = "video " + segment.location
                    segment.locator["kind"] = "video"
                    result.append(segment)
        return result
    raise ValueError(f"Unsupported file type: {extension}")
