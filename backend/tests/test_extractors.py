from pathlib import Path
import pytest
from docx import Document
from pptx import Presentation
from reportlab.pdfgen.canvas import Canvas
from PIL import Image, ImageDraw, ImageFont
import pandas as pd
from app.services.extract import extract
from app.services.extract import Segment
from app.rag.chunking import split


def test_office_pdf_tables_and_image_ocr(tmp_path: Path):
    pdf = tmp_path / 'notice.pdf'
    page = Canvas(str(pdf)); page.drawString(80, 700, 'Cooling pump inspection scheduled Friday'); page.save()
    assert 'Cooling pump inspection' in ' '.join(s.text for s in extract(pdf, '.pdf'))

    document = Document(); document.add_paragraph('Maintenance owner is Priya')
    table = document.add_table(rows=2, cols=2); table.cell(0,0).text='Asset'; table.cell(0,1).text='Cost'; table.cell(1,0).text='Pump'; table.cell(1,1).text='500'
    docx = tmp_path / 'plan.docx'; document.save(docx)
    parts = extract(docx, '.docx')
    assert any('Maintenance owner' in s.text for s in parts)
    assert any(s.modality == 'table' and 'Pump' in s.text for s in parts)

    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[5]); slide.shapes.title.text='Incident timeline'
    pptx = tmp_path / 'deck.pptx'; deck.save(pptx)
    assert any('Incident timeline' in s.text for s in extract(pptx, '.pptx'))

    xlsx = tmp_path / 'metrics.xlsx'; pd.DataFrame({'day':['Mon','Tue'], 'pressure':[10,15]}).to_excel(xlsx,index=False)
    assert any(s.modality == 'table' and 'pressure' in s.text for s in extract(xlsx, '.xlsx'))

    image = Image.new('RGB', (800, 180), 'white')
    font = ImageFont.load_default(size=48)
    ImageDraw.Draw(image).text((35, 45), 'PRESSURE 42', fill='black', font=font)
    png = tmp_path / 'chart.png'; image.save(png)
    assert 'PRESSURE 42' in extract(png, '.png')[0].text.upper()


def test_pdf_chunk_anchors_follow_the_split_passage():
    first = 'A' * 100
    second = 'Cooling inspection belongs to the second paragraph and has a distinct source anchor.'
    part = Segment(first + '\n\n' + second, 'page 2', locator={'kind': 'pdf', 'page': 2, 'anchor': first})
    chunks = split([part], method='recursive', size=100)
    assert len(chunks) == 2
    assert chunks[0].text == first
    assert chunks[1].text == second
    assert chunks[1].locator['anchor'].startswith('Cooling inspection')


def test_large_preserved_source_is_fully_chunked():
    source = 'head\n' + ('A' * 24000) + 'LAST_EVIDENCE'
    chunks = split([Segment(source, 'JSON', 'logs', preserve=True, locator={'kind': 'json'})])
    assert len(chunks) > 1
    assert ''.join(chunk.text for chunk in chunks) == source
    assert all(len(chunk.text) <= 8000 for chunk in chunks)
    assert all('oversized_segment_split' in chunk.quality['flags'] for chunk in chunks)
    assert chunks[-1].locator['part_count'] == len(chunks)


def test_large_table_keeps_rows_and_repeats_header():
    header = 'asset,value\n'
    source = header + ''.join(f'asset-{i},{i}\n' for i in range(1000))
    chunks = split([Segment(source, 'sheet Main', 'table', preserve=True)])
    assert len(chunks) > 1
    assert all(len(chunk.text) <= 8000 for chunk in chunks)
    assert all(chunk.text.startswith(header) for chunk in chunks)
    assert chunks[0].text + ''.join(chunk.text[len(header):] for chunk in chunks[1:]) == source
    assert all(chunk.text.endswith('\n') for chunk in chunks)


def test_csv_row_limit_fails_instead_of_silently_omitting_evidence(tmp_path: Path):
    csv = tmp_path / 'records.csv'
    csv.write_text('id,value\n' + ''.join(f'{i},ok\n' for i in range(100001)))
    with pytest.raises(ValueError, match='100000-row indexing limit'):
        extract(csv, '.csv')
